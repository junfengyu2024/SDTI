#!/usr/bin/env python
# coding=utf-8
# Copyright 2025 Yixiao Chen, Shikun Li, Ruoyu Li, Zhe Li, Junliang Xing
# Copyright 2024 The HuggingFace Inc. team. All rights reserved.
# Licensed under the Apache License, Version 2.0
"""
Inference script for multi‑targeted attack: frozen first‑half denoising +
second‑half generation, attacking standard and robust models.
Target labels are loaded from the NIPS CSV and are guaranteed to differ
from the true class.
"""

import os
import json
import argparse

import pandas as pd
import safetensors
import torch
import torch.nn.functional as F
import torch.utils.data
from accelerate.utils import set_seed
from peft import LoraConfig
from peft.utils import set_peft_model_state_dict
from PIL import Image
from torchvision import transforms
from tqdm.auto import tqdm
from transformers import CLIPTextModel, CLIPTokenizer

import lpips
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from diffusers import (
    AutoencoderKL,
    UNet2DConditionModel,
    DDIMScheduler,
    DDIMInverseScheduler,
)
from diffusers.utils import convert_all_state_dict_to_peft

from utils import (
    get_classes,
    normalize,
    unnormalize_ddpm,
    budget,
    get_all_models,
    load_robust_model,
    fix_labels_nips,
)

IMG_EXTENSIONS = (".jpg", ".jpeg", ".png", ".ppm", ".bmp", ".pgm", ".tif", ".tiff", ".webp")

ROBUST_MODEL_STATS = {
    'res50_sin':         (224, (0.485, 0.456, 0.406), (0.229, 0.224, 0.225)),
    'res50_sin_in':      (224, (0.485, 0.456, 0.406), (0.229, 0.224, 0.225)),
    'res50_sin_fine_in': (224, (0.485, 0.456, 0.406), (0.229, 0.224, 0.225)),
    'res50_augmix':      (224, (0.485, 0.456, 0.406), (0.229, 0.224, 0.225)),
    'adv_incv3':         (299, (0.5, 0.5, 0.5), (0.5, 0.5, 0.5)),
    'ens_inc_res_v2':    (299, (0.5, 0.5, 0.5), (0.5, 0.5, 0.5)),
}


class SafeImageDataset(torch.utils.data.Dataset):
    """Recursively loads images skipping hidden files / folders.
       When target_labels is provided, each sample returns (image, target_label)."""
    def __init__(self, root_dir, transform=None, target_labels=None):
        self.image_paths = []
        for cur_root, dirs, files in os.walk(root_dir):
            dirs[:] = [d for d in dirs if not d.startswith('.')]
            for f in files:
                if not f.startswith('.') and f.lower().endswith(IMG_EXTENSIONS):
                    self.image_paths.append(os.path.join(cur_root, f))
        self.image_paths.sort()
        if not self.image_paths:
            raise FileNotFoundError(f"No valid images found in {root_dir}")
        self.transform = transform
        self.target_labels = target_labels
        if target_labels is not None and len(target_labels) != len(self.image_paths):
            raise ValueError("Length of target_labels must match number of images.")

    def __len__(self):
        return len(self.image_paths)

    def __getitem__(self, idx):
        img = Image.open(self.image_paths[idx]).convert("RGB")
        if self.transform:
            img = self.transform(img)
        label = self.target_labels[idx] if self.target_labels is not None else 0
        return img, label


def normalize_with_stats(x, mean, std):
    mean_t = torch.tensor(mean, device=x.device, dtype=x.dtype).view(1, -1, 1, 1)
    std_t = torch.tensor(std, device=x.device, dtype=x.dtype).view(1, -1, 1, 1)
    return (x - mean_t) / std_t


def save_results_to_excel(excel_path, model_names, asrs, lpips_val, metadata):
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Model", "ASR", "LPIPS"])

    hfill = PatternFill("solid", fgColor="1F4E78")
    hfont = Font(color="FFFFFF", bold=True)
    border = Border(bottom=Side(style="thin", color="BFBFBF"))
    for col in range(1, 4):
        cell = ws.cell(row=1, column=col)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for name in model_names:
        ws.append([name, asrs.get(name, 0.0), lpips_val])

    for row in ws.iter_rows(min_row=2):
        row[1].number_format = '0.00%'
        row[2].number_format = '0.0000'

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 14
    ws.freeze_panes = "B2"

    ws2 = wb.create_sheet("Metadata")
    ws2.append(["Key", "Value"])
    ws2["A1"].fill = hfill
    ws2["B1"].fill = hfill
    ws2["A1"].font = hfont
    ws2["B1"].font = hfont
    for k, v in metadata.items():
        ws2.append([k, str(v)])
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 56
    wb.save(excel_path)


def parse_args():
    parser = argparse.ArgumentParser(description="Evaluate multi‑targeted adversarial attack.")
    parser.add_argument("--pretrained_model_name_or_path", type=str, required=True)
    parser.add_argument("--dataset_dir", type=str, required=True)
    parser.add_argument("--lora_path", type=str, required=True)
    parser.add_argument("--label_to_caption_path", type=str, required=True)
    parser.add_argument("--excel_output_path", type=str, required=True)
    parser.add_argument("--images_csv", type=str, default="images.csv",
                        help="CSV file (inside dataset_dir) with ImageId and TargetClass columns.")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--test_batch_size", type=int, default=16)
    parser.add_argument("--dataloader_num_workers", type=int, default=8)
    parser.add_argument("--eps", type=int, default=16)
    parser.add_argument("--label_flag", type=str, default="N8")
    parser.add_argument("--model_type", type=str, required=True)
    parser.add_argument("--s_for", type=int, default=6)
    parser.add_argument("--s_gen", type=int, default=6)
    parser.add_argument("--num_timesteps", type=int, default=20)
    parser.add_argument("--freeze_ratio", type=float, default=0.5)
    parser.add_argument("--rank", type=int, default=16)
    parser.add_argument("--include_standard_models", action="store_true", default=True)
    parser.add_argument("--robust_models_list", nargs="*",
                        default=['res50_sin', 'res50_sin_in', 'res50_sin_fine_in',
                                 'res50_augmix', 'adv_incv3', 'ens_inc_res_v2'])
    return parser.parse_args()


def main():
    args = parse_args()
    torch.set_float32_matmul_precision("high")

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    weight_dtype = torch.float32
    eps = args.eps / 255.0

    set_seed(args.seed)

    scale_size = 299 if args.model_type == "incv3" else 224

    # --------------- load target labels from NIPS CSV ---------------
    csv_path = os.path.join(args.dataset_dir, args.images_csv)
    csv_df = pd.read_csv(csv_path)
    valid_ids = set(str(x) for x in csv_df["ImageId"])

    # gather all image paths, keep only those present in CSV
    all_image_paths = []
    for root, dirs, files in os.walk(args.dataset_dir):
        dirs[:] = [d for d in dirs if not d.startswith('.')]
        for f in files:
            if not f.startswith('.') and f.lower().endswith(IMG_EXTENSIONS):
                stem = os.path.splitext(f)[0]
                if stem in valid_ids:
                    all_image_paths.append(os.path.join(root, f))
    all_image_paths.sort()

    # dummy dataset to feed fix_labels_nips
    class DummySet:
        pass
    dummy = DummySet()
    dummy.samples = [(p, 0) for p in all_image_paths]

    target_set = fix_labels_nips(dummy, args.dataset_dir, pytorch=True, target_flag=True)
    target_labels = [label for _, label in target_set.samples]

    # --------------- diffusion setup ---------------
    tokenizer = CLIPTokenizer.from_pretrained(args.pretrained_model_name_or_path, subfolder="tokenizer")
    text_encoder = CLIPTextModel.from_pretrained(args.pretrained_model_name_or_path, subfolder="text_encoder")
    vae = AutoencoderKL.from_pretrained(args.pretrained_model_name_or_path, subfolder="vae")
    unet = UNet2DConditionModel.from_pretrained(args.pretrained_model_name_or_path, subfolder="unet")
    unet_origin = UNet2DConditionModel.from_pretrained(args.pretrained_model_name_or_path, subfolder="unet")
    ddim_scheduler = DDIMScheduler.from_pretrained(args.pretrained_model_name_or_path, subfolder="scheduler")
    ddim_inv_scheduler = DDIMInverseScheduler.from_pretrained(args.pretrained_model_name_or_path, subfolder="scheduler")

    ddim_inv_scheduler.set_timesteps(args.num_timesteps, device=device)
    timesteps_inverse = ddim_inv_scheduler.timesteps[:args.s_for]
    ddim_scheduler.set_timesteps(args.num_timesteps, device=device)
    timesteps = ddim_scheduler.timesteps[-args.s_gen:]

    freeze_steps = max(1, int(len(timesteps) * args.freeze_ratio))
    freeze_steps = min(freeze_steps, len(timesteps) - 1)
    timesteps_frozen = timesteps[:freeze_steps]
    timesteps_train = timesteps[freeze_steps:]

    print(f"Timesteps: {len(timesteps)} total, {len(timesteps_frozen)} frozen, {len(timesteps_train)} trainable")

    lpips_fn = lpips.LPIPS(net="alex").to(device, dtype=weight_dtype).eval()
    lpips_fn.requires_grad_(False)

    def freeze(model):
        model.requires_grad_(False)
        return model.to(device, dtype=weight_dtype)

    vae = freeze(vae)
    text_encoder = freeze(text_encoder)
    unet_origin = freeze(unet_origin)
    unet = unet.to(device, dtype=weight_dtype)

    # LoRA
    unet_lora_config = LoraConfig(
        r=args.rank, lora_alpha=args.rank, init_lora_weights="gaussian",
        target_modules=["to_k", "to_q", "to_v", "to_out.0", "add_k_proj", "add_v_proj"],
    )
    unet.add_adapter(unet_lora_config)

    lora_state_dict = safetensors.torch.load_file(args.lora_path)
    unet_lora_state_dict = {
        k[len("unet."):]: v for k, v in lora_state_dict.items() if k.startswith("unet.")
    }
    unet_lora_state_dict = convert_all_state_dict_to_peft(unet_lora_state_dict)
    set_peft_model_state_dict(unet, unet_lora_state_dict)
    print("LoRA loaded.")
    unet.eval().requires_grad_(False)

    def tokenize(captions):
        return tokenizer(captions, max_length=tokenizer.model_max_length,
                         padding="max_length", truncation=True, return_tensors="pt").input_ids

    # --------------- dataset ---------------
    test_transform = transforms.Compose([
        transforms.Resize(scale_size),
        transforms.CenterCrop(scale_size),
        transforms.ToTensor(),
        transforms.Normalize([0.5] * 3, [0.5] * 3),
    ])

    base_dataset = SafeImageDataset(args.dataset_dir, test_transform, target_labels=target_labels)
    print(f"Test images: {len(base_dataset)}")

    with open(args.label_to_caption_path, "r") as f:
        label_to_caption = json.load(f)

    fixed_labels = target_labels
    fixed_captions = [f"a photo of {label_to_caption[str(l)][1]}" for l in fixed_labels]

    class TargetDataset(torch.utils.data.Dataset):
        def __init__(self, base, labels, captions):
            self.base = base
            self.labels = labels
            self.captions = captions
        def __len__(self): return len(self.base)
        def __getitem__(self, idx):
            img, _ = self.base[idx]
            return img, self.labels[idx], self.captions[idx]

    test_dataset = TargetDataset(base_dataset, fixed_labels, fixed_captions)

    def collate(batch):
        imgs = torch.stack([b[0] for b in batch]).float()
        labels = torch.tensor([b[1] for b in batch], dtype=torch.long)
        captions = [b[2] for b in batch]
        return {"imgs": imgs, "labels": labels, "captions": tokenize(captions)}

    test_loader = torch.utils.data.DataLoader(
        test_dataset, batch_size=args.test_batch_size,
        shuffle=False, collate_fn=collate, num_workers=args.dataloader_num_workers,
    )

    # --------------- victim models ---------------
    attacked = []
    if args.include_standard_models:
        for cfg in get_all_models(device, weight_dtype):
            cfg["source"] = "standard"
            attacked.append(cfg)

    for name in args.robust_models_list:
        if name not in ROBUST_MODEL_STATS:
            continue
        try:
            model = load_robust_model(name, device).to(dtype=weight_dtype).eval()
        except Exception as e:
            print(f"Skip robust model {name}: {e}")
            continue
        in_sz, mean, std = ROBUST_MODEL_STATS[name]
        attacked.append({"name": name, "model": model, "input_size": in_sz,
                         "mean": mean, "std": std, "source": "robust"})

    print(f"Victim models: {len(attacked)}")

    # --------------- evaluation ---------------
    unet_in_transform = transforms.Resize(352)
    unet_out_transform = transforms.Resize(scale_size)

    asr_success = {m["name"]: 0 for m in attacked}
    total_lpips = 0.0
    total_imgs = 0

    for batch in tqdm(test_loader, desc="Attacking"):
        imgs = batch["imgs"].to(dtype=weight_dtype, device=device)
        labels = batch["labels"].to(device)
        cap = batch["captions"].to(device)

        with torch.no_grad():
            latents = vae.encode(unet_in_transform(imgs)).latent_dist.sample() * vae.config.scaling_factor
            enc = text_encoder(cap)[0]
            enc_empty = text_encoder(tokenize([""] * imgs.size(0)).to(device))[0]

            for t in timesteps_inverse:
                pred = unet_origin(latents, t, encoder_hidden_states=enc_empty)[0]
                latents = ddim_inv_scheduler.step(pred, t, latents)[0]

            for t in timesteps_frozen:
                pred = unet_origin(latents, t, encoder_hidden_states=enc_empty)[0]
                latents, _ = ddim_scheduler.step(pred, t, latents)

            for t in timesteps_train:
                pred = unet(latents, t, encoder_hidden_states=enc)[0]
                latents, x0_pred = ddim_scheduler.step(pred, t, latents)

            adv = vae.decode(x0_pred / vae.config.scaling_factor)[0]
            adv = budget(unet_out_transform(unnormalize_ddpm(adv)),
                         unnormalize_ddpm(imgs), eps, device, attack_mode="multi_targeted")

        clean = torch.clamp(unnormalize_ddpm(imgs), 0, 1)
        adv = torch.clamp(adv, 0, 1)

        lpips_vals = lpips_fn(clean, adv).view(-1)
        bs = imgs.size(0)
        total_lpips += lpips_vals.sum().item()
        total_imgs += bs

        for mcfg in attacked:
            model = mcfg["model"]
            inp_size = mcfg["input_size"]
            inp = F.interpolate(adv, size=inp_size, mode="bilinear", align_corners=False) if inp_size != scale_size else adv

            if mcfg["source"] == "robust":
                inp = normalize_with_stats(inp, mcfg["mean"], mcfg["std"])
            else:
                inp = normalize(inp)

            out = model(inp)
            logits = out.logits if hasattr(out, "logits") else out
            preds = logits.argmax(1)
            asr_success[mcfg["name"]] += (preds == labels).sum().item()

    lpips_avg = total_lpips / total_imgs if total_imgs else 0
    asr_final = {k: v / total_imgs if total_imgs else 0 for k, v in asr_success.items()}

    metadata = {k: str(getattr(args, k)) for k in [
        "dataset_dir", "eps", "s_for", "s_gen", "num_timesteps", "freeze_ratio", "lora_path"
    ]}
    metadata["robust_models"] = args.robust_models_list
    metadata["num_images"] = total_imgs

    save_results_to_excel(args.excel_output_path, list(asr_final.keys()), asr_final, lpips_avg, metadata)
    print(f"Results saved to {args.excel_output_path}")


if __name__ == "__main__":
    main()
